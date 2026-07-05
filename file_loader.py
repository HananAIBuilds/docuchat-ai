"""
file_loader.py
--------------
Extracts raw text from uploaded documents.

Supports: .txt, .pdf, .docx, .csv, .xlsx

Design notes:
- Works directly on in-memory bytes (io.BytesIO), so it plugs straight into
  Streamlit's `st.file_uploader` without ever touching disk.
- Each loader is isolated so adding a new file type later is a one-function change.
"""

import io
import csv

SUPPORTED_EXTENSIONS = [".txt", ".pdf", ".docx", ".csv", ".xlsx"]


def get_extension(filename: str) -> str:
    if "." not in filename:
        return ""
    return "." + filename.lower().rsplit(".", 1)[-1]


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Route to the correct extractor based on file extension.

    Raises:
        ValueError: if the file type isn't supported.
    """
    ext = get_extension(filename)

    if ext == ".txt":
        return _read_txt(file_bytes)
    elif ext == ".pdf":
        return _read_pdf(file_bytes)
    elif ext == ".docx":
        return _read_docx(file_bytes)
    elif ext == ".csv":
        return _read_csv(file_bytes)
    elif ext == ".xlsx":
        return _read_xlsx(file_bytes)
    else:
        raise ValueError(
            f"Unsupported file type '{ext or filename}'. "
            f"Supported types: {', '.join(SUPPORTED_EXTENSIONS)}"
        )


def _read_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore")


def _read_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts)


def _read_docx(file_bytes: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_csv(file_bytes: bytes) -> str:
    decoded = file_bytes.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(decoded))
    rows = [", ".join(row) for row in reader]
    return "\n".join(rows)


def _read_xlsx(file_bytes: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    text = "\n".join(
        ", ".join(str(cell) if cell is not None else "" for cell in row) for row in rows
    )
    workbook.close()
    return text
